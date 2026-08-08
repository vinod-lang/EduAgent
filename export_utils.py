import io
from docx import Document as DocxDocument
from fpdf import FPDF


def generate_docx_bytes(text):
    """
    Converts plain text (like a generated notice) into a real .docx
    file, kept in memory so it can be offered as a download without
    ever writing to disk.
    """
    doc = DocxDocument()
    for line in text.split("\n"):
        doc.add_paragraph(line)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def generate_quiz_pdf_bytes(questions, title="Quiz"):
    """
    Builds a real PDF containing the quiz on one page and the
    answer key on a separate page — exactly how a professor would
    want to print/distribute it.
    """
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 10, title, ln=True)
    pdf.set_font("Helvetica", size=12)

    for i, q in enumerate(questions, start=1):
        pdf.multi_cell(0, 8, f"Q{i}. {q['question']}")
        if "options" in q:
            for letter, opt in q["options"].items():
                pdf.multi_cell(0, 8, f"   {letter}) {opt}")
        pdf.ln(2)

    # Answer key on a fresh page, separated from the questions
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 14)
    pdf.cell(0, 10, "Answer Key", ln=True)
    pdf.set_font("Helvetica", size=12)

    for i, q in enumerate(questions, start=1):
        if "correct_answer" in q:
            pdf.multi_cell(0, 8, f"Q{i}: {q['correct_answer']} - {q.get('explanation', '')}")
        else:
            pdf.multi_cell(0, 8, f"Q{i}: {q.get('model_answer', '')}")

    pdf_bytes = bytes(pdf.output())
    return io.BytesIO(pdf_bytes)