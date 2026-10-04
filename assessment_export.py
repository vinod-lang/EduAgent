"""Separate student paper and suggested answer key; exports stay in memory."""
import io
import os
from pathlib import Path
from assessment_spec import AssessmentError

class AssessmentExportError(AssessmentError):
    pass


def export_blocks(result,answer_key=False):
    spec=result.spec
    blocks=[]
    if spec.institution:blocks.append(('institution',spec.institution))
    blocks.append(('title',spec.title+(' Answer Key' if answer_key else '')))
    blocks.append(('body',f'{spec.assessment_type} | {spec.scope.course} | {spec.scope.subject} | {spec.scope.semester}'))
    blocks.append(('body',f'Total marks: {spec.total_marks} | Questions: {spec.total_questions}'))
    if answer_key:
        blocks.append(('body','Professor reference only. Descriptive answers are suggested model answers; review before grading.'))
    else:blocks.append(('body',spec.instructions))
    previous=None
    for q in result.questions:
        if q.question_type!=previous:
            blocks.append(('section','Multiple Choice Questions' if q.question_type=='MCQ' else 'Descriptive Questions'))
            previous=q.question_type
        blocks.append(('question',f'Q{q.question_number}. [{q.marks} marks] {q.question_text}'))
        if answer_key:
            if q.question_type=='MCQ':blocks.append(('body',f'Correct answer: {q.correct_answer} - {dict(q.options)[q.correct_answer]}'))
            blocks.append(('body','Suggested answer: '+q.model_answer))
        else:
            for letter,text in q.options:blocks.append(('option',f'{letter}) {text}'))
    return blocks


def _font(text):
    from fontTools.ttLib import TTFont
    override=os.environ.get('EDUAGENT_ASSESSMENT_FONT')
    candidates=[Path(override)] if override else [Path('/System/Library/Fonts/Supplemental/Arial Unicode.ttf'),Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')]
    for path in candidates:
        if not path.is_file():continue
        try:
            with TTFont(path,lazy=True) as font:
                glyphs=font.getBestCmap() or {}
                if all(c.isspace() or ord(c) in glyphs for c in text):return str(path)
        except Exception:
            continue
    raise AssessmentExportError('No available assessment font covers this text. Set EDUAGENT_ASSESSMENT_FONT to a suitable local Unicode TTF font; no text was silently dropped.')


def assessment_pdf_bytes(result,*,answer_key=False):
    from fpdf import FPDF
    blocks=export_blocks(result,answer_key);font=_font(''.join(text for _,text in blocks))
    class Paper(FPDF):
        def footer(self):
            self.set_y(-12);self.set_font('Assessment',size=9)
            self.cell(0,6,f'Page {self.page_no()}',align='R')
    pdf=Paper();pdf.set_margins(14,14,14);pdf.set_auto_page_break(True,margin=18)
    try:
        pdf.add_font('Assessment',fname=font);pdf.add_page()
        pdf.set_title(result.spec.title);pdf.set_author('EduAgent')
        for kind,text in blocks:
            size={'title':18,'institution':11,'section':13,'question':11,'body':10,'option':10}[kind]
            if kind in ('section','question') and pdf.get_y()>pdf.h-40:pdf.add_page()
            pdf.set_x(pdf.l_margin);pdf.set_font('Assessment',size=size)
            if kind=='section':pdf.ln(3)
            pdf.multi_cell(0,7,text,align='C' if kind in ('title','institution') else 'L',new_x='LMARGIN',new_y='NEXT')
            pdf.ln(2 if kind!='title' else 5)
        return bytes(pdf.output())
    except Exception as exc:
        raise AssessmentExportError('Assessment PDF could not be formatted; reduce unusually long content and retry.') from exc


def assessment_docx_bytes(result,*,answer_key=False):
    from docx import Document
    from docx.shared import Inches,Pt,RGBColor
    doc=Document();section=doc.sections[0]
    section.top_margin=section.bottom_margin=Inches(.7)
    style=doc.styles['Normal'];style.font.name='Arial';style.font.size=Pt(11)
    doc.styles['Title'].font.color.rgb=RGBColor(0,0,0)
    for kind,text in export_blocks(result,answer_key):
        if kind=='title':doc.add_paragraph(text,'Title')
        elif kind=='section':doc.add_heading(text,level=1)
        else:
            p=doc.add_paragraph(text)
            if kind=='question':p.paragraph_format.keep_with_next=True
            if kind=='option':p.paragraph_format.left_indent=Inches(.2)
    buffer=io.BytesIO();doc.save(buffer);return buffer.getvalue()
