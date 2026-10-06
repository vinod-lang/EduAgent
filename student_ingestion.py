"""Bounded, in-memory CSV/XLSX parsing and deterministic student normalization.

No storage, models, formula evaluation, or analytics calculations live here.
"""
from dataclasses import dataclass
import csv
import io
import math
from pathlib import Path
import re
import unicodedata
import zipfile

import pandas as pd

MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 10_000
MAX_COLUMNS = 100
MAX_SHEETS = 20
MAX_CELLS = 1_000_000
MAX_EXPANDED_BYTES = 50 * 1024 * 1024
MAX_CELL_CHARS = 10_000
MISSING = {'', 'na', 'n/a', 'null', 'none', '-', 'nan'}


class StudentDataError(ValueError):
    """A controlled file or mapping error safe to display to the professor."""


@dataclass(frozen=True)
class Sheet:
    name: str
    hidden: bool
    rows: tuple


@dataclass(frozen=True)
class Workbook:
    sheets: tuple
    kind: str

    def sheet(self, name):
        match = next((s for s in self.sheets if s.name == name), None)
        if match is None:
            raise StudentDataError('Select an available sheet.')
        return match


@dataclass(frozen=True)
class Issue:
    row: int
    column: str
    code: str
    severity: str
    message: str


@dataclass
class RawTable:
    frame: pd.DataFrame
    source_rows: tuple
    original_headers: tuple
    issues: tuple


@dataclass(frozen=True)
class Suggestion:
    candidates: tuple
    confidence: str

    @property
    def selected(self):
        return self.candidates[0] if len(self.candidates) == 1 else None


@dataclass(frozen=True)
class Assessment:
    column: str
    name: str
    scale: str = 'percentage'
    maximum: float | None = None


@dataclass(frozen=True)
class Mapping:
    student_id: str | None
    student_name: str | None
    attendance: str
    assessments: tuple
    attendance_scale: str = 'auto'
    assessment_number: str | None = None


@dataclass
class NormalizedDataset:
    frame: pd.DataFrame
    issues: tuple
    row_states: pd.DataFrame
    summary: dict
    mode: str

    def validation_frame(self):
        return pd.DataFrame([vars(i) for i in self.issues], columns=['row','column','code','severity','message'])


def label_key(value):
    return re.sub(r'[^\w]+', '', unicodedata.normalize('NFKC', str(value)).casefold()).replace('_', '')


ALIASES = {
    'student_id': {'rollno','rollnumber','enrollmentno','enrollmentnumber','enrolmentno','studentid','registrationno','registrationnumber'},
    'student_name': {'name','studentname','nameofstudent'},
    'attendance': {'attendance','attendancepercent','attendancepercentage','attendancepct'},
    'assessment_number': {'assessmentnumber','assessmentno'},
}
MARK_WORDS = ('quiz','midsem','midterm','endsem','assignment','total','marks','score','final','test','exam')


def missing(value):
    if value is None:
        return True
    if isinstance(value, str):
        return value.strip().casefold() in MISSING
    return bool(pd.isna(value))


def _bounded_rows(rows):
    result = []
    cells = 0
    for row in rows:
        if len(result) >= MAX_ROWS + 20:
            raise StudentDataError(f'Sheet exceeds {MAX_ROWS} student rows plus 20 header/title rows.')
        if len(row) > MAX_COLUMNS:
            raise StudentDataError(f'Sheet exceeds {MAX_COLUMNS} columns.')
        cells += len(row)
        if cells > MAX_CELLS:
            raise StudentDataError('Sheet exceeds the cell limit.')
        if any(isinstance(v, str) and len(v) > MAX_CELL_CHARS for v in row):
            raise StudentDataError('A cell exceeds the text length limit.')
        result.append(tuple(row))
    return tuple(result)


def parse_student_file(data, filename):
    """Accept bytes only. XLSX is read-only, with formulas exposed, never evaluated."""
    suffix = Path(filename).suffix.lower()
    if suffix not in {'.csv', '.xlsx'}:
        raise StudentDataError('Supported student files are CSV and XLSX; XLS/XLSM are not supported.')
    if not isinstance(data, bytes) or not data or not data.strip():
        raise StudentDataError('Student file is empty.')
    if len(data) > MAX_BYTES:
        raise StudentDataError('Student file exceeds the 10 MB limit.')
    if suffix == '.csv':
        try:
            text = data.decode('utf-8-sig')
            if '\x00' in text:
                raise StudentDataError('CSV contains invalid null characters.')
            sample = text[:65536]
            try:
                delimiter = csv.Sniffer().sniff(sample, delimiters=',;\t|').delimiter
            except csv.Error:
                # The first nonblank header supplies a conservative fallback.
                lines = [r for r in text.splitlines()[:20] if r.strip()]
                candidates = [d for d in ',;\t|' if any(d in line for line in lines)]
                if len(candidates) != 1:
                    raise StudentDataError('Cannot determine CSV delimiter; use UTF-8 comma-separated CSV.')
                delimiter = candidates[0]
            rows = _bounded_rows(csv.reader(io.StringIO(text), delimiter=delimiter, strict=True))
        except StudentDataError:
            raise
        except (UnicodeError, csv.Error) as exc:
            raise StudentDataError('Cannot read CSV. Use valid UTF-8 text with consistent quoting.') from exc
        if not any(any(not missing(v) for v in row) for row in rows):
            raise StudentDataError('CSV contains no data.')
        return Workbook((Sheet('CSV', False, rows),), 'csv')
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(e.file_size for e in entries) > MAX_EXPANDED_BYTES:
                raise StudentDataError('Workbook expanded contents exceed the safe resource limit.')
            if any(e.flag_bits & 1 for e in entries):
                raise StudentDataError('Encrypted workbooks are unsupported.')
            if any('vbaproject' in e.filename.casefold() for e in entries):
                raise StudentDataError('Macro-containing workbooks are unsupported.')
        import openpyxl
        from openpyxl.xml import DEFUSEDXML
        if not DEFUSEDXML:
            raise StudentDataError('Safe Excel XML parser is unavailable; install the pinned dependencies.')
        book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        try:
            if len(book.worksheets) > MAX_SHEETS:
                raise StudentDataError(f'Workbook exceeds {MAX_SHEETS} sheets.')
            sheets = []
            total_cells = 0
            for ws in book.worksheets:
                if ws.max_row and ws.max_row > MAX_ROWS + 20 or ws.max_column and ws.max_column > MAX_COLUMNS:
                    raise StudentDataError('Workbook sheet dimensions exceed the row/column limits.')
                rows = []
                for row in ws.iter_rows():
                    values = []
                    for cell in row:
                        value = cell.value
                        # Preserve simple zero-padded identifier display, without formula execution.
                        if cell.data_type != 'f' and isinstance(value, int) and re.fullmatch('0{2,}', cell.number_format or ''):
                            value = str(value).zfill(len(cell.number_format))
                        values.append(value)
                    rows.append(values)
                    if len(rows) > MAX_ROWS + 20:
                        raise StudentDataError('Sheet row limit exceeded.')
                bounded = _bounded_rows(rows)
                total_cells += sum(len(row) for row in bounded)
                if total_cells > MAX_CELLS:
                    raise StudentDataError('Workbook exceeds the total cell limit.')
                sheets.append(Sheet(ws.title, ws.sheet_state != 'visible', bounded))
            return Workbook(tuple(sheets), 'xlsx')
        finally:
            book.close()
    except StudentDataError:
        raise
    except Exception as exc:
        # Parser internals/filenames/raw cells do not leak into error text.
        raise StudentDataError('Cannot read XLSX. Choose a valid, unencrypted workbook.') from exc


def header_candidates(sheet):
    """Advisory scores only; the professor always chooses/confirms a row."""
    scores = []
    for number, row in enumerate(sheet.rows[:20], 1):
        keys = [label_key(v) for v in row if not missing(v)]
        score = sum(3 for aliases in ALIASES.values() if any(k in aliases for k in keys))
        score += sum(1 for k in keys if any(word in k for word in MARK_WORDS))
        if score:
            scores.append((number, score))
    scores.sort(key=lambda v: (-v[1], v[0]))
    if not scores:
        return (), 'Unmapped'
    best = tuple(row for row, score in scores if score == scores[0][1])
    return best, 'Ambiguous' if len(best) > 1 else 'High confidence' if scores[0][1] >= 7 else 'Possible match'


def make_raw_table(sheet, header_row):
    if not isinstance(header_row, int) or isinstance(header_row, bool) or not 1 <= header_row <= len(sheet.rows):
        raise StudentDataError('Choose a header row present in the sheet.')
    raw_header = sheet.rows[header_row - 1]
    if not raw_header or all(missing(v) for v in raw_header):
        raise StudentDataError('Selected header row is empty.')
    headers = []; originals = []; issues = []
    for i, value in enumerate(raw_header, 1):
        original = '' if missing(value) else str(value).strip()
        name = original or f'Unnamed column {i}'
        if name in headers:
            name = f'{name} [column {i}]'
        while name in headers:
            name += '_'
        if name != original:
            issues.append(Issue(header_row, name, 'header_disambiguated', 'WARNING', 'Blank/duplicate header has a distinct display label; review mapping.'))
        headers.append(name); originals.append(original)
    rows = []; positions = []
    for number, row in enumerate(sheet.rows[header_row:], header_row + 1):
        if all(missing(v) for v in row):
            continue
        if len(row) > len(headers) and any(not missing(v) for v in row[len(headers):]):
            raise StudentDataError(f'Row {number} has more values than the selected header.')
        rows.append(list(row[:len(headers)]) + [None] * max(0, len(headers) - len(row)))
        positions.append(number)
    if not rows:
        raise StudentDataError('Selected sheet/header contains no student rows.')
    if len(rows) > MAX_ROWS:
        raise StudentDataError(f'Dataset exceeds {MAX_ROWS} rows.')
    return RawTable(pd.DataFrame(rows, columns=headers), tuple(positions), tuple(originals), tuple(issues))


def suggest_columns(table):
    suggestions = {}
    for role, aliases in ALIASES.items():
        candidates = tuple(c for c, original in zip(table.frame, table.original_headers) if label_key(original) in aliases)
        suggestions[role] = Suggestion(candidates, 'Ambiguous' if len(candidates) > 1 else 'High confidence' if candidates else 'Unmapped')
    reserved = {c for s in suggestions.values() for c in s.candidates}
    candidates = tuple(c for c, original in zip(table.frame, table.original_headers)
                       if c not in reserved and any(w in label_key(original) for w in MARK_WORDS)
                       and not table.frame[c].map(missing).all())
    suggestions['assessments'] = Suggestion(candidates, 'Possible match' if candidates else 'Unmapped')
    return suggestions


def suggested_maximum(label):
    match = re.search(r'\((\d+(?:\.\d+)?)\)|\bout\s+of\s+(\d+(?:\.\d+)?)', str(label), re.I)
    return float(next(v for v in match.groups() if v is not None)) if match else None


def _number(value):
    if isinstance(value, bool):
        raise StudentDataError('Boolean values are not academic numbers.')
    try:
        result = float(str(value).strip())
    except (TypeError, ValueError):
        raise StudentDataError('Expected a numeric value.') from None
    if not math.isfinite(result):
        raise StudentDataError('Expected a finite numeric value.')
    return result


def normalize_attendance(value, scale='auto'):
    if scale not in {'auto', 'percentage', 'fraction'}:
        raise StudentDataError('Unknown attendance scale.')
    if missing(value):
        return float('nan')
    text = str(value).strip()
    explicit_percent = text.endswith('%')
    if explicit_percent:
        if scale == 'fraction':
            raise StudentDataError('Percent notation conflicts with fraction mode.')
        result = _number(text[:-1])
    else:
        result = _number(value)
        if scale == 'fraction':
            if not 0 <= result <= 1:
                raise StudentDataError('Fraction attendance must be 0–1.')
            result *= 100
        elif scale == 'auto':
            if 0 < result < 1:
                result *= 100
            elif 1 <= result < 2:
                raise StudentDataError('Ambiguous attendance: choose percentage or fraction mode explicitly.')
    if not 0 <= result <= 100:
        raise StudentDataError('Attendance must be within 0–100% after normalization.')
    return result


def normalize_mark(value, assessment):
    if missing(value):
        return float('nan')
    text = str(value).strip()
    if '/' in text:
        parts = text.split('/')
        if len(parts) != 2:
            raise StudentDataError('Expected a single score/maximum pair.')
        score, denominator = map(_number, parts)
        if denominator <= 0 or not 0 <= score <= denominator:
            raise StudentDataError('Score must be within 0–maximum.')
        if assessment.scale == 'raw' and denominator != assessment.maximum:
            raise StudentDataError('Score denominator does not match the configured maximum.')
        if assessment.scale != 'raw':
            raise StudentDataError('Score/maximum notation requires raw-mark mode and a configured maximum.')
        return score / denominator * 100
    if text.endswith('%'):
        if assessment.scale != 'percentage':
            raise StudentDataError('Percent notation conflicts with raw-mark mode.')
        result = _number(text[:-1])
    else:
        result = _number(value)
    maximum = 100 if assessment.scale == 'percentage' else assessment.maximum
    if maximum is None or not math.isfinite(maximum) or maximum <= 0:
        raise StudentDataError('Raw marks require a finite positive maximum.')
    if not 0 <= result <= maximum:
        raise StudentDataError('Mark is outside the configured range; values are not clamped.')
    return result / maximum * 100


def _validate_mapping(table, mapping):
    fields = [mapping.student_id, mapping.student_name, mapping.attendance, mapping.assessment_number]
    fields = [v for v in fields if v is not None]
    if not mapping.student_id and not mapping.student_name:
        raise StudentDataError('Map at least Student ID or Student Name.')
    if mapping.attendance_scale not in {'auto','percentage','fraction'}:
        raise StudentDataError('Choose a valid attendance scale.')
    if not mapping.assessments or not mapping.attendance:
        raise StudentDataError('Map attendance and at least one assessment column.')
    names = []
    for a in mapping.assessments:
        if not isinstance(a.name, str) or not a.name.strip() or len(a.name) > 100:
            raise StudentDataError('Assessment names must be nonblank and at most 100 characters.')
        if a.name.strip() in {'student_id','student_name','attendance','attendance_percent','assessment_number'}:
            raise StudentDataError('Assessment name conflicts with a canonical identity/attendance field.')
        if a.scale not in {'raw','percentage'}:
            raise StudentDataError('Choose raw or percentage marks.')
        if a.scale == 'raw' and (not isinstance(a.maximum,(int,float)) or isinstance(a.maximum,bool) or not math.isfinite(a.maximum) or a.maximum <= 0):
            raise StudentDataError('Raw marks require a finite positive maximum.')
        fields.append(a.column); names.append(a.name.strip())
    if len(set(fields)) != len(fields) or any(v not in table.frame for v in fields):
        raise StudentDataError('Mapped fields must be distinct, existing source columns.')
    if len(set(names)) != len(names):
        raise StudentDataError('Assessment display names must be unique.')
    if mapping.assessment_number and len(mapping.assessments) != 1:
        raise StudentDataError('History mode requires one score column plus an assessment-number column.')


def normalize_dataset(table, mapping):
    _validate_mapping(table, mapping)
    mode = 'trend' if mapping.assessment_number else 'snapshot'
    issues = list(table.issues); rows = []; states = []; identities = []
    for pos, (_, source) in enumerate(table.frame.iterrows()):
        number = table.source_rows[pos]; start = len(issues)
        def add(column, code, severity, message):
            issues.append(Issue(number, column, code, severity, message))
        ident = None if not mapping.student_id or missing(source[mapping.student_id]) else str(source[mapping.student_id]).strip()
        name = None if not mapping.student_name or missing(source[mapping.student_name]) else str(source[mapping.student_name]).strip()
        # Only exact summary labels, with no competing identity, are excluded.
        summaries = {'average','total','class average','grand total'}
        if (name or ident or '').casefold() in summaries and not (ident and name and ident.casefold() not in summaries):
            add(mapping.student_name or mapping.student_id, 'summary_row', 'INFO', 'Explicit summary row excluded from student analysis.')
            states.append({'row':number,'state':'excluded'}); continue
        if mapping.student_id and not ident:
            add(mapping.student_id,'missing_id','ERROR','Student ID is missing; mixed identity rows require correction.')
        if mapping.student_name and not name:
            add(mapping.student_name,'missing_name','ERROR','Student Name is missing.')
        display_name = name or ident or ''  # ID-only inputs display the ID, never an invented name.
        identity = ident if mapping.student_id else name
        row = {'student_name':display_name}
        if mapping.student_id:
            row['student_id'] = ident
        for column in (mapping.student_id, mapping.student_name):
            if column and isinstance(source[column], str) and source[column] != source[column].strip():
                add(column,'identity_trimmed','INFO','Leading/trailing identity whitespace cleaned.')
            if column and not missing(source[column]) and str(source[column]).startswith('='):
                add(column,'formula_identity','ERROR','Formula identity is unsupported; provide literal identifiers/names.')
        try:
            row['attendance_percent'] = normalize_attendance(source[mapping.attendance], mapping.attendance_scale)
            if missing(source[mapping.attendance]):
                add(mapping.attendance,'missing_attendance','WARNING','Attendance is missing; record remains incomplete.')
            elif mapping.attendance_scale == 'auto' and not str(source[mapping.attendance]).endswith('%') and 0 < _number(source[mapping.attendance]) < 1:
                add(mapping.attendance,'fraction_interpreted','INFO','Fractional attendance interpreted as 0–1; review the selected scale.')
        except StudentDataError as exc:
            row['attendance_percent'] = float('nan')
            add(mapping.attendance,'invalid_attendance','ERROR',str(exc))
        for a in mapping.assessments:
            target = 'marks' if mode == 'trend' else a.name.strip()
            try:
                row[target] = normalize_mark(source[a.column], a)
                if missing(source[a.column]):
                    add(a.column,'missing_mark','WARNING','Assessment mark is missing; not converted to zero.')
            except StudentDataError as exc:
                row[target] = float('nan'); add(a.column,'invalid_mark','ERROR',str(exc))
        if mode == 'trend':
            try:
                value = _number(source[mapping.assessment_number])
                if value <= 0 or not value.is_integer():
                    raise StudentDataError('Assessment number must be a positive integer.')
                row['assessment_number'] = value
            except StudentDataError as exc:
                row['assessment_number'] = None; add(mapping.assessment_number,'invalid_assessment_number','ERROR',str(exc))
        row_issues = issues[start:]
        state = 'invalid' if any(i.severity=='ERROR' for i in row_issues) else 'incomplete' if any(i.code.startswith('missing_') for i in row_issues) else 'valid'
        rows.append((number,row)); states.append({'row':number,'state':state})
        identities.append((number,identity,row.get('assessment_number'),display_name))
    groups = {}
    for number, identity, assessment, name in identities:
        if identity is not None:
            key = (identity, assessment) if mode == 'trend' else identity
            groups.setdefault(key,[]).append(number)
    duplicate_rows = {number for numbers in groups.values() if len(numbers)>1 for number in numbers}
    conflicting = set()
    if mode == 'trend' and mapping.student_id:
        names = {}
        for number, identity, assessment, name in identities:
            names.setdefault(identity,set()).add(name)
        conflicting = {number for number,identity,_,_ in identities if identity and len(names[identity])>1}
    for state in states:
        number = state['row']
        if number in duplicate_rows:
            code = 'duplicate_id' if mapping.student_id else 'duplicate_identity'
            issues.append(Issue(number,mapping.student_id or mapping.student_name,code,'ERROR','Duplicate identity/assessment requires professor review; no rows merged.'))
            state['state'] = 'invalid'
        if number in conflicting:
            issues.append(Issue(number,mapping.student_name,'inconsistent_name','ERROR','Student ID has inconsistent names across assessments.'))
            state['state'] = 'invalid'
    if mode == 'trend':
        invalid_rows = {s['row'] for s in states if s['state']=='invalid'}
        invalid_ids = {identity for number,identity,_,_ in identities if number in invalid_rows and identity}
        row_identity = {number:identity for number,identity,_,_ in identities}
        for state in states:
            identity = row_identity.get(state['row'])
            if identity in invalid_ids and state['state'] not in {'invalid','excluded'}:
                state['state'] = 'invalid'
                issues.append(Issue(state['row'], mapping.student_id or mapping.student_name, 'invalid_history', 'ERROR', 'Another row for this identity is invalid; full history excluded until corrected.'))
    eligible = {s['row'] for s in states if s['state'] in {'valid','incomplete'}}
    frame = pd.DataFrame([r for number,r in rows if number in eligible])
    if frame.empty:
        columns = ['student_name','attendance_percent'] + (['student_id'] if mapping.student_id else []) + (['marks','assessment_number'] if mode=='trend' else [a.name.strip() for a in mapping.assessments])
        frame = pd.DataFrame(columns=columns)
    state_frame = pd.DataFrame(states,columns=['row','state'])
    error_rows = {i.row for i in issues if i.severity=='ERROR'}
    warning_rows = {i.row for i in issues if i.severity=='WARNING' and i.row in table.source_rows}
    summary = dict(total_rows=len(table.frame), valid_rows=sum(s['state']=='valid' for s in states), incomplete_rows=sum(s['state']=='incomplete' for s in states), invalid_rows=sum(s['state']=='invalid' for s in states), excluded_rows=sum(s['state']=='excluded' for s in states), eligible_rows=len(frame), rows_with_errors=len(error_rows), rows_with_warnings=len(warning_rows), missing_values=sum(i.code.startswith('missing_') for i in issues), duplicate_ids=sum(i.code=='duplicate_id' for i in issues), invalid_attendance=sum(i.code=='invalid_attendance' for i in issues), invalid_marks=sum(i.code=='invalid_mark' for i in issues))
    return NormalizedDataset(frame,tuple(issues),state_frame,summary,mode)
