import io
from dataclasses import replace
import pdfplumber
import pytest
from docx import Document
from assessment_studio import validate_output
from assessment_export import assessment_pdf_bytes,assessment_docx_bytes,AssessmentExportError
from assessment_fixtures import spec,raw,evidence,ID_A

def result():return validate_output(raw(),spec(),evidence())

@pytest.mark.parametrize('answer_key',[False,True])
def test_pdf_separate_answers_marks_provenance(answer_key):
    data=assessment_pdf_bytes(result(),answer_key=answer_key)
    assert data.startswith(b'%PDF')
    with pdfplumber.open(io.BytesIO(data)) as pdf:text='\n'.join(p.extract_text() or '' for p in pdf.pages)
    assert 'Total marks: 5' in text and 'Q1.' in text and '3 marks' in text and '2 marks' in text
    assert ('SUGGESTED_ANSWER' in text)==answer_key
    assert ('A) Maximum variance' in text)==(not answer_key)
    assert ID_A not in text and 'owned_chunk' not in text and 'Synthetic lecture.pdf' not in text and '0.2' not in text
    if answer_key:assert 'Correct answer: A' in text and 'suggested model answers' in text

@pytest.mark.parametrize('answer_key',[False,True])
def test_docx_separate_answers(answer_key):
    doc=Document(io.BytesIO(assessment_docx_bytes(result(),answer_key=answer_key)))
    text='\n'.join(p.text for p in doc.paragraphs)
    assert ('SUGGESTED_ANSWER' in text)==answer_key
    assert 'Total marks: 5' in text and ID_A not in text and 'Synthetic lecture.pdf' not in text
    assert any(p.style.name=='Title' for p in doc.paragraphs)

@pytest.mark.parametrize('format',['PDF','DOCX'])
def test_unicode_survives(format):
    original=result();updated=replace(original,spec=replace(original.spec,title='講義 Δ Assessment'))
    if format=='PDF':
        with pdfplumber.open(io.BytesIO(assessment_pdf_bytes(updated))) as pdf:text='\n'.join(p.extract_text() or '' for p in pdf.pages)
    else:text='\n'.join(p.text for p in Document(io.BytesIO(assessment_docx_bytes(updated))).paragraphs)
    assert '講義 Δ Assessment' in text

def test_missing_unicode_font_controlled(monkeypatch):
    monkeypatch.setenv('EDUAGENT_ASSESSMENT_FONT','/nonexistent.ttf')
    with pytest.raises(AssessmentExportError,match='font'):assessment_pdf_bytes(result())


def test_long_paper_wraps_across_pages():
    import json
    from assessment_fixtures import output
    s=spec(question_types={'MCQ':4,'Descriptive':6},difficulties={'Easy':2,'Medium':4,'Hard':4},
           blooms=dict(Remember=2,Understand=2,Apply=2,Analyze=2,Evaluate=1,Create=1),total_questions=10,total_marks=50)
    data=output(s.plan)
    for q in data['questions']:
        q['question_text']+=' '+('Explain the variance of the projection using the supplied academic context. '*8)
    paper=validate_output(json.dumps(data),s,evidence())
    with pdfplumber.open(io.BytesIO(assessment_pdf_bytes(paper))) as pdf:
        assert len(pdf.pages)>1
        text='\n'.join(page.extract_text() or '' for page in pdf.pages)
        assert 'Q10.' in text and '50' in text and 'SUGGESTED_ANSWER' not in text
        for page in pdf.pages:
            assert all(0<=c['x0']<c['x1']<=page.width+.1 and 0<=c['top']<c['bottom']<=page.height+.1 for c in page.chars)
