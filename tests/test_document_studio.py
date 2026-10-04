import io
import json
from dataclasses import replace
from unittest.mock import Mock
import pytest
import pdfplumber
from docx import Document
import db
import document_studio as studio
from document_models import *
from document_templates import *
from document_export import *
from document_repository import *
from ai_provider import AIConnectionError


def draft(kind='permission_request',**kw):
    return DocumentDraft(kind,title='Workshop permission',recipient='Head of Department',subject='Workshop request',body=('Please permit a workshop on 15 November 2026 for 80 students with a budget of ₹25,000.',),**kw)

@pytest.mark.parametrize('kind',CATALOG)
def test_request_catalog(kind):
    request=DocumentRequest(kind,'  Synthetic professor instructions  ')
    assert request.description=='Synthetic professor instructions' and request.date==''

@pytest.mark.parametrize('kwargs',[{'description':''},{'description':' '},{'document_type':'fake'},{'tone':'fake'},{'template_id':'../private'},{'description':'x'*20001},{'description':None},{'description':'bad\x00'}])
def test_request_invalid(kwargs):
    values=dict(document_type='notice',description='Synthetic');values.update(kwargs)
    with pytest.raises(DocumentError):DocumentRequest(**values)

@pytest.mark.parametrize('tone',TONES)
def test_tones_unicode_optional(tone):
    assert DocumentRequest('notice','Draft a ₹25,000 notice 講義',tone,recipient=' Students ').recipient=='Students'


def test_generate_structured_payload(monkeypatch):
    request=DocumentRequest('permission_request','Request a workshop on 15 November 2026',recipient='Head of Department',title='Workshop permission',subject='Workshop request')
    spy=Mock(return_value=draft().to_json());monkeypatch.setattr(studio.ai_provider,'generate_chat',spy)
    assert studio.generate_draft(request)==draft()
    payload=json.loads(spy.call_args.kwargs['messages'][1]['content'])
    assert payload['description']==request.description and payload['type_structure']=='letter'
    assert 'Do not invent' in spy.call_args.kwargs['messages'][0]['content']

@pytest.mark.parametrize('change',[{'body':[]},{'body':'prose'},{'body':[{}]},{'body':['']},{'document_type':'fake'},{'debug':'secret'},{'body':[None]}])
def test_invalid_model_schema(change):
    data=draft().to_dict();data.update(change)
    with pytest.raises(DocumentError):parse_draft(json.dumps(data))

@pytest.mark.parametrize('raw',['bad','```json {} ```','[]','null','{"body":[],"body":[]}','{"body":NaN}'])
def test_bad_json(raw):
    with pytest.raises(DocumentError):parse_draft(raw)

@pytest.mark.parametrize('field',['body','title','recipient','subject'])
def test_missing_fields(field):
    data=draft().to_dict();del data[field]
    with pytest.raises(DocumentError):parse_draft(json.dumps(data))

@pytest.mark.parametrize('field',['recipient','sender','date','reference_number','signature'])
def test_no_invented_optional_metadata(field):
    data=DocumentDraft('notice',body=('Synthetic notice',)).to_dict();data[field]='Invented'
    with pytest.raises(DocumentError):parse_draft(json.dumps(data),DocumentRequest('notice','Synthetic notice'))


def test_provider_failure(monkeypatch):
    monkeypatch.setattr(studio.ai_provider,'generate_chat',Mock(side_effect=AIConnectionError('Synthetic unavailable')))
    with pytest.raises(AIConnectionError):studio.generate_draft(DocumentRequest('notice','Synthetic'))

@pytest.mark.parametrize('kind',['notice','official_email','official_letter','financial_approval','permission_request','memo','custom'])
def test_type_structure(kind):
    style=CATALOG[kind][1]
    d=DocumentDraft(kind,title='Synthetic heading',recipient='Synthetic audience',subject='Synthetic purpose',body=('Synthetic body',),salutation='Dear recipient,' if style in ('email','letter') else '',closing='Regards,' if style in ('email','letter') else '')
    blocks=STANDARD.blocks(d)
    assert ('salutation','Dear recipient,') in blocks if style in ('email','letter') else not any(k=='salutation' for k,_ in blocks)
    assert not any(k=='title' for k,_ in blocks) if style=='email' else ('title','Synthetic heading') in blocks
    if style in ('announcement','memo','report'):
        with pytest.raises(DocumentError):replace(d,salutation='Dear recipient')

@pytest.mark.parametrize('placeholders',[('body','{{eval(1)}}'),('body','__class__'),('title',),('body','body')])
def test_unsafe_template(placeholders):
    with pytest.raises(DocumentError):DocumentTemplate('local','Local',placeholders)


def test_template_optional_unicode_long():
    d=DocumentDraft('notice',body=('講義 Δ '+ 'Synthetic '*900,))
    assert STANDARD.blocks(d)==(('body',d.body[0]),)
    template=DocumentTemplate('future','Future',('date','body','signature'))
    assert template.blocks(d)==(('body',d.body[0]),)


def test_versions_and_exports_current():
    original=draft();versions=DocumentVersions.generated(original)
    edited=replace(original,title='Edited heading',subject='Edited subject',recipient='Edited recipient',body=('Edited body content',),closing='Respectfully,',signature='Synthetic Professor')
    versions=versions.update(edited)
    assert versions.original==original and versions.current==edited and versions.undo().current==original
    doc=Document(io.BytesIO(document_docx_bytes(versions.current)))
    docx='\n'.join(p.text for p in doc.paragraphs)
    with pdfplumber.open(io.BytesIO(document_pdf_bytes(versions.current))) as pdf: pdftext='\n'.join(p.extract_text() or '' for p in pdf.pages)
    for text in ('Edited heading','Edited subject','Edited recipient','Edited body content','Respectfully,','Synthetic Professor'):
        assert text in docx and text in pdftext
    assert original.body[0] not in docx and original.body[0] not in pdftext
    assert any(p.style.name=='Title' for p in doc.paragraphs)
    assert doc.sections[0].left_margin.inches==pytest.approx(.85)
    assert 'document_type' not in docx and 'standard_academic' not in pdftext

@pytest.mark.parametrize('format',['pdf','docx'])
def test_long_unicode_export(format):
    d=DocumentDraft('report_submission',title='講義 Δ Report',body=tuple('Paragraph '+str(i)+' '+('Synthetic readable material '*50) for i in range(20)))
    if format=='pdf':
        with pdfplumber.open(io.BytesIO(document_pdf_bytes(d))) as pdf:
            assert len(pdf.pages)>1
            text='\n'.join(p.extract_text() or '' for p in pdf.pages)
            assert '講義 Δ Report' in text and 'Paragraph 19' in text
            for page in pdf.pages: assert all(0<=c['x0']<c['x1']<=page.width+.1 and 0<=c['top']<c['bottom']<=page.height+.1 for c in page.chars)
    else:
        doc=Document(io.BytesIO(document_docx_bytes(d)))
        assert '講義 Δ Report' in doc.paragraphs[0].text and 'Paragraph 19' in doc.paragraphs[-1].text


def test_missing_font_controlled(monkeypatch):
    monkeypatch.setenv('EDUAGENT_DOCUMENT_FONT','/nonexistent.ttf')
    with pytest.raises(DocumentExportError):document_pdf_bytes(draft())


def test_refinement_current_preserved(monkeypatch):
    current=draft();short=replace(current,body=('Permission for 80 students on 15 November 2026, budget ₹25,000.',))
    spy=Mock(return_value=short.to_json());monkeypatch.setattr(studio.ai_provider,'generate_chat',spy)
    versions=DocumentVersions.generated(current)
    updated=versions.update(studio.refine_draft(versions.current,'Make this shorter.'))
    assert updated.original==current and updated.current==short and updated.undo().current==current
    messages=spy.call_args.kwargs['messages'];payload=json.loads(messages[1]['content'])
    assert payload['current']['body']==list(current.body) and payload['instruction']=='Make this shorter.'
    assert 'Preserve names, dates, amounts' in messages[0]['content']
    assert 'must exactly match the optional request fields' not in messages[0]['content']

@pytest.mark.parametrize('response',['bad',draft(date='Invented').to_json(),replace(draft(),body=('Changed date 2027',)).to_json()])
def test_refinement_invalid_keeps_previous(monkeypatch,response):
    versions=DocumentVersions.generated(draft());monkeypatch.setattr(studio.ai_provider,'generate_chat',Mock(return_value=response))
    with pytest.raises(DocumentError):studio.refine_draft(versions.current,'Make it shorter')
    assert versions.current==versions.original and len(versions.history)==1

@pytest.fixture
def isolated_db(tmp_path,monkeypatch):
    monkeypatch.setattr(db,'DB_PATH',str(tmp_path/'documents.db'));db.init_db();return tmp_path


def test_persistence_additive_idempotent_legacy(isolated_db):
    db.add_document_record('Legacy material','C','U1','legacy.pdf')
    before=db.get_documents_for_course('C')
    with db.material_connection() as conn:
        init_document_schema(conn);init_document_schema(conn)
    versions=DocumentVersions.generated(draft())
    identity=save_draft(versions);assert load_draft(identity)==versions
    edited=versions.update(replace(draft(),body=('Edited persisted body',)))
    save_draft(edited,identity,'Final');loaded=load_draft(identity)
    assert loaded.original==versions.original and loaded.current==edited.current
    assert len(list_drafts())==1 and list_drafts()[0]['status']=='Final'
    assert db.get_documents_for_course('C')==before
    assert all(row['details']=='' for row in db.get_recent_activity() if row['action']=='document_saved')


def test_empty_schema_and_read_no_migration(isolated_db):
    assert list_drafts()==[]
    with db.material_connection() as conn:
        assert not conn.execute("select 1 from sqlite_master where name='document_drafts'").fetchone()
        init_document_schema(conn)
    assert list_drafts()==[]


def test_missing_saved_record_controlled(isolated_db):
    with pytest.raises(DocumentStorageError):save_draft(DocumentVersions.generated(draft()),'missing')

@pytest.mark.parametrize('value',[None,[],{},42])
def test_invalid_document_type_controlled(value):
    with pytest.raises(DocumentError):DocumentDraft(value,body=('Synthetic',))
    with pytest.raises(DocumentError):DocumentRequest(value,'Synthetic')


def test_template_never_evaluates_content():
    body='{{__import__("os").system("invalid")}}'
    assert STANDARD.blocks(DocumentDraft('custom',body=(body,)))==(('body',body),)


def test_all_required_structural_fields():
    for field in draft().to_dict():
        data=draft().to_dict();del data[field]
        with pytest.raises(DocumentError):parse_draft(json.dumps(data))


def test_persistence_different_original_rejected(isolated_db):
    original=DocumentVersions.generated(draft());identity=save_draft(original)
    unrelated=DocumentVersions.generated(replace(draft(),body=('Different original',)))
    with pytest.raises(DocumentStorageError):save_draft(unrelated,identity)
    assert load_draft(identity).original==original.original


def test_refinement_do_not_change_is_not_fact_change_permission(monkeypatch):
    changed=replace(draft(),body=('Workshop in 2027',))
    monkeypatch.setattr(studio.ai_provider,'generate_chat',Mock(return_value=changed.to_json()))
    with pytest.raises(DocumentError):studio.refine_draft(draft(),'Make this shorter; do not change dates or amounts.')


def test_refinement_explicit_fact_change(monkeypatch):
    changed=replace(draft(),body=('Workshop in 2027',))
    monkeypatch.setattr(studio.ai_provider,'generate_chat',Mock(return_value=changed.to_json()))
    assert studio.refine_draft(draft(),'Change the workshop date to 2027.')==changed
