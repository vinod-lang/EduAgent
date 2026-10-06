"""Build 23 synthetic session-owned assessment review/edit/export contracts."""
import json
from unittest.mock import Mock
from dataclasses import replace
import pytest
import ai_provider
from test_api import web,client
from test_security import secured,services,upload,own
from security.models import Scope
from assessment_spec import BLOOMS

BODY=dict(assessment_type='Quiz',course='ML',semester='S1',subject='ML',units=['Unit 1'],material_ids=[],question_types={'MCQ':1,'Descriptive':1},difficulties={'Easy':1,'Medium':1,'Hard':0},blooms={k:int(k in ('Remember','Understand')) for k in BLOOMS},total_questions=2,total_marks=5,title='Synthetic assessment')
def raw_questions():
    return {'questions':[dict(question_number=i,question_type='MCQ' if i==1 else 'Descriptive',question_text=f'Synthetic PCA question {i}?',options={'A':'Projection','B':'Deletion','C':'Storage','D':'Attendance'} if i==1 else {},correct_answer='A' if i==1 else '',model_answer='PCA projects data.',difficulty='Easy' if i==1 else 'Medium',bloom_level='Remember' if i==1 else 'Understand',marks=3 if i==1 else 2,evidence_ids=['E1']) for i in (1,2)]}
@pytest.fixture
def assessment(web,services,secured,monkeypatch):
    a=upload(services[0],secured,secured.a,'Synthetic PCA')
    provider=Mock(return_value=json.dumps(raw_questions()));monkeypatch.setattr(ai_provider,'generate_chat',provider)
    c=client(web);body=dict(BODY,material_ids=[a['material_id']]);r=c.post('/api/v1/assessments/generate',json=body);assert r.status_code==200,r.text
    return c,r.json(),provider,body

def edit_body(a):
    return {'revision':a['revision'],'questions':[{k:q[k] for k in ('question_text','marks','options','correct_answer','model_answer')} for q in a['questions']]}

def test_review_and_safe_validation_summary(assessment):
    c,a,_,_=assessment;r=c.get('/api/v1/assessments/'+a['handle']);assert r.json()==a
    assert all(a['validation'].values()) and not a['professor_edited'];assert all(k not in r.text for k in ('evidence_ids','teaching_evidence','pyq_text','response_hash'))

@pytest.mark.parametrize('operation',['read','edit','export','discard'])
def test_workspace_cross_professor_and_cross_session_isolation(assessment,web,operation):
    _,a,_,_=assessment;path='/api/v1/assessments/'+a['handle']
    for alias in ('b','a'):
        other=client(web,alias)
        r=other.get(path) if operation=='read' else other.patch(path,json=edit_body(a)) if operation=='edit' else other.get(path+'/export') if operation=='export' else other.delete(path)
        assert r.status_code==404

def test_edits_revalidate_without_ai_and_export_current_content(assessment):
    c,a,provider,_=assessment;body=edit_body(a);body['questions'][0]['question_text']='Professor revised PCA question?';body['questions'][0]['marks']=2;body['questions'][1]['marks']=3
    r=c.patch('/api/v1/assessments/'+a['handle'],json=body);assert r.status_code==200,r.text;assert r.json()['revision']==1 and r.json()['professor_edited'];assert r.json()['questions'][0]['marks']==2;assert provider.call_count==1
    import io
    from docx import Document
    paper=c.get('/api/v1/assessments/'+a['handle']+'/export?format=docx');text='\n'.join(p.text for p in Document(io.BytesIO(paper.content)).paragraphs);assert 'Professor revised PCA question?' in text and 'Correct answer:' not in text
    key=c.get('/api/v1/assessments/'+a['handle']+'/export?format=docx&answer_key=true');assert 'Correct answer:' in '\n'.join(p.text for p in Document(io.BytesIO(key.content)).paragraphs)
    assert c.get('/api/v1/assessments/'+a['handle']+'/export').content.startswith(b'%PDF')

@pytest.mark.parametrize('change',[('marks',0),('marks',2),('marks',True),('question_text',''),('options',{'A':'same','B':'same','C':'same','D':'same'}),('options',{'A':'Only'}),('correct_answer','Z'),('model_answer','')])
def test_invalid_edits_preserve_previous_valid_workspace(assessment,change):
    c,a,provider,_=assessment;body=edit_body(a);body['questions'][0][change[0]]=change[1];path='/api/v1/assessments/'+a['handle'];r=c.patch(path,json=body);assert r.status_code==422;assert c.get(path).json()==a;assert provider.call_count==1;assert 'Traceback' not in r.text

def test_unknown_fields_cannot_replace_server_evidence_or_type(assessment):
    c,a,_,_=assessment;body=edit_body(a);body['questions'][0]['evidence_ids']=['injected'];assert c.patch('/api/v1/assessments/'+a['handle'],json=body).status_code==422

def test_stale_revision_is_rejected(assessment):
    c,a,_,_=assessment;path='/api/v1/assessments/'+a['handle'];body=edit_body(a);assert c.patch(path,json=body).status_code==200;assert c.patch(path,json=body).status_code==409

def test_count_changes_are_rejected(assessment):
    c,a,_,_=assessment;body=edit_body(a);body['questions'].pop();assert c.patch('/api/v1/assessments/'+a['handle'],json=body).status_code==422

def test_discard_and_logout_invalidate_workspaces(assessment):
    c,a,_,body=assessment;path='/api/v1/assessments/'+a['handle'];assert c.delete(path).status_code==200;assert c.get(path).status_code==404
    another=c.post('/api/v1/assessments/generate',json=body).json()['handle'];assert c.post('/api/v1/auth/logout').status_code==200;assert c.get('/api/v1/assessments/'+another).status_code==401

def test_denied_private_and_guessed_materials_never_call_provider(web,services,secured,monkeypatch):
    other=upload(services[0],secured,secured.b,'Private B');provider=Mock();monkeypatch.setattr(ai_provider,'generate_chat',provider);c=client(web)
    for identity in (other['material_id'],'11111111-1111-4111-8111-111111111111'):
        assert c.post('/api/v1/assessments/generate',json=dict(BODY,material_ids=[identity])).status_code==404
    provider.assert_not_called()

def test_course_membership_generation_is_enforced(web,services,secured,monkeypatch):
    shared=upload(services[0],secured,secured.a,'shared',own(secured.a,Scope.COURSE,secured.course));provider=Mock(return_value=json.dumps(raw_questions()));monkeypatch.setattr(ai_provider,'generate_chat',provider);c=client(web,'b');body=dict(BODY,material_ids=[shared['material_id']]);assert c.post('/api/v1/assessments/generate',json=body).status_code==404;provider.assert_not_called();secured.repo.set_membership(secured.b.professor_id,secured.course);assert c.post('/api/v1/assessments/generate',json=body).status_code==200

def test_rejected_generation_diagnostic_is_safe_and_no_assessment_stored(web,services,secured,monkeypatch):
    material=upload(services[0],secured,secured.a,'safe');monkeypatch.setattr(ai_provider,'generate_chat',Mock(return_value='RAW PRIVATE /path invalid response'));c=client(web);r=c.post('/api/v1/assessments/generate',json=dict(BODY,material_ids=[material['material_id']]));assert r.status_code==422;assert r.json()['error']['diagnostic']['category']=='INVALID_JSON';assert 'RAW PRIVATE' not in r.text and '/path' not in r.text;assert not any(row[1]=='assessment' for row in web.state.workspaces.rows.values())

def test_pyq_is_session_owned_and_still_protected_on_edit(web,services,secured,monkeypatch):
    item=upload(services[0],secured,secured.a,'PYQ teaching');monkeypatch.setattr(services[0].assessments,'extract_pyq',lambda *args,**kw:'Explain a synthetic separating margin in detail?');provider=Mock(return_value=json.dumps(raw_questions()));monkeypatch.setattr(ai_provider,'generate_chat',provider);c=client(web);handle=c.post('/api/v1/assessments/pyq',files={'file':('synthetic.pdf',b'synthetic')}).json()['handle']
    body=dict(BODY,material_ids=[item['material_id']],pyq_handle=handle)
    assert client(web,'a').post('/api/v1/assessments/generate',json=body).status_code==404;provider.assert_not_called()
    a=c.post('/api/v1/assessments/generate',json=body).json();edited=edit_body(a);edited['questions'][0]['question_text']='Explain a synthetic separating margin in detail?'
    assert c.patch('/api/v1/assessments/'+a['handle'],json=edited).status_code==422
    assert c.get('/api/v1/assessments/'+a['handle']).json()==a
