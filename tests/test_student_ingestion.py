"""Synthetic parser/mapping/normalization fixtures; no production student data."""
import io
import math
import zipfile
from unittest.mock import Mock
import pandas as pd
import pytest
import openpyxl
import student_ingestion as ingestion
from student_ingestion import *
from student_hub import *


def workbook_bytes(sheets=None):
    book=openpyxl.Workbook();book.remove(book.active)
    for name,rows,hidden in sheets or [('Section A',[['Roll No','Name','Attendance','Quiz (20)'],['001','Student A','82%',16]],False)]:
        ws=book.create_sheet(name)
        for row in rows:ws.append(row)
        if hidden:ws.sheet_state='hidden'
    stream=io.BytesIO();book.save(stream);book.close();return stream.getvalue()


def raw(rows=None,headers=('Roll No','Name','Attendance','Quiz')):
    rows=rows or [('001','Student A',90,80)]
    return make_raw_table(Sheet('Synthetic',False,(headers,*rows)),1)


def mapping(**changes):
    values=dict(student_id='Roll No',student_name='Name',attendance='Attendance',assessments=(Assessment('Quiz','Quiz'),),attendance_scale='percentage')
    values.update(changes);return Mapping(**values)

@pytest.mark.parametrize('delimiter',[',',';','\t','|'])
@pytest.mark.parametrize('bom',[False,True])
def test_csv_formats(delimiter,bom):
    text=delimiter.join(['Roll No','Name','Attendance','Quiz'])+'\n'+delimiter.join(['001','Student A','90','80'])+'\n'
    data=text.encode('utf-8-sig' if bom else 'utf-8')
    table=make_raw_table(parse_student_file(data,'students.CSV').sheets[0],1)
    assert table.frame.iloc[0,0]=='001' and len(table.frame)==1


def test_csv_blanks_headers_duplicate_empty_columns():
    book=parse_student_file(b' Roll No , Name , Attendance , Quiz , Quiz ,\n001,A,90,80,70,\n,,,,,\n','x.csv')
    table=make_raw_table(book.sheets[0],1)
    assert table.frame.columns.tolist()==['Roll No','Name','Attendance','Quiz','Quiz [column 5]','Unnamed column 6']
    assert len(table.frame)==1 and len(table.issues)==2

@pytest.mark.parametrize('filename,data,match',[('x.xls',b'file','Supported'),('x.xlsm',b'file','Supported'),('x.csv',b'','empty'),('x.csv',b' \n','empty'),('x.csv',b'\xff,abc','UTF-8'),('x.csv',b'A,B\n"unclosed,2','CSV'),('x.xlsx',b'broken','XLSX'),('x.csv',b'\0,A','null'),('x.csv',b'singlecolumn\nvalue','delimiter')])
def test_controlled_bad_files(filename,data,match):
    with pytest.raises(StudentDataError,match=match):parse_student_file(data,filename)


def test_file_size_limit(monkeypatch):
    monkeypatch.setattr(ingestion,'MAX_BYTES',3)
    with pytest.raises(StudentDataError,match='limit'):parse_student_file(b'A,B\n','x.csv')


def test_extra_csv_values_rejected():
    sheet=parse_student_file(b'A,B\n1,2,3\n','x.csv').sheets[0]
    with pytest.raises(StudentDataError,match='more values'):make_raw_table(sheet,1)


def test_xlsx_single_sheet():
    book=parse_student_file(workbook_bytes(),'synthetic.xlsx')
    assert book.kind=='xlsx' and len(book.sheets)==1
    table=make_raw_table(book.sheets[0],1)
    normalized=normalize_dataset(table,mapping(assessments=(Assessment('Quiz (20)','Quiz','raw',20),)))
    assert normalized.frame.student_id.tolist()==['001']
    assert normalized.frame.Quiz.tolist()==[80] and normalized.frame.attendance_percent.tolist()==[82]


def test_xlsx_multisheet_empty_hidden_unicode():
    book=parse_student_file(workbook_bytes([('Section A',[['Name','Attendance','Quiz'],['Synthetic',90,80]],False),('空のシート',[],False),('Hidden',[['Name','Attendance','Quiz'],['学生',90,80]],True)]),'x.xlsx')
    assert [s.name for s in book.sheets]==['Section A','空のシート','Hidden']
    assert book.sheet('Hidden').hidden
    table=make_raw_table(book.sheet('Hidden'),1)
    assert table.frame.Name.tolist()==['学生']
    with pytest.raises(StudentDataError):book.sheet('Missing')
    with pytest.raises(StudentDataError):make_raw_table(book.sheet('空のシート'),1)


def test_formulas_not_evaluated():
    book=parse_student_file(workbook_bytes([('Data',[['Roll No','Name','Attendance','Quiz'],['001','A','=100/2','=2+3']],False)]),'x.xlsx')
    table=make_raw_table(book.sheets[0],1)
    assert table.frame.Quiz.iloc[0]=='=2+3'
    result=normalize_dataset(table,mapping())
    assert result.frame.empty and result.summary['invalid_marks']==1 and result.summary['invalid_attendance']==1


def test_excel_padded_identifier():
    book=openpyxl.Workbook();ws=book.active
    ws.append(['Roll No','Name','Attendance','Quiz']);ws.append([1,'Synthetic',90,80]);ws['A2'].number_format='000'
    stream=io.BytesIO();book.save(stream);book.close()
    assert make_raw_table(parse_student_file(stream.getvalue(),'x.xlsx').sheets[0],1).frame.iloc[0,0]=='001'

@pytest.mark.parametrize('limit,value,match',[('MAX_COLUMNS',3,'dimensions'),('MAX_ROWS',1,'dimensions'),('MAX_SHEETS',0,'sheets'),('MAX_CELLS',2,'cell')])
def test_workbook_resource_limits(monkeypatch,limit,value,match):
    data=workbook_bytes([('Data',[['Name','Attendance','Quiz']]+[['Synthetic',90,80]]*22,False)]) if limit=='MAX_ROWS' else workbook_bytes();monkeypatch.setattr(ingestion,limit,value)
    with pytest.raises(StudentDataError,match=match):parse_student_file(data,'x.xlsx')


def test_zip_expansion_limit(monkeypatch):
    data=workbook_bytes();monkeypatch.setattr(ingestion,'MAX_EXPANDED_BYTES',10)
    with pytest.raises(StudentDataError,match='expanded'):parse_student_file(data,'x.xlsx')


def test_macro_rejected():
    stream=io.BytesIO()
    with zipfile.ZipFile(stream,'w') as book:book.writestr('xl/vbaProject.bin',b'synthetic')
    with pytest.raises(StudentDataError,match='Macro'):parse_student_file(stream.getvalue(),'x.xlsx')


def test_title_blank_header_detection_and_override():
    sheet=Sheet('Data',False,(('Synthetic results',),(),('Roll No','Name','Attendance','Quiz'),('001','A',90,80)))
    candidates,confidence=header_candidates(sheet)
    assert candidates==(3,) and confidence=='High confidence'
    assert make_raw_table(sheet,3).source_rows==(4,)
    with pytest.raises(StudentDataError,match='more values'):make_raw_table(sheet,1)


def test_ambiguous_header_rows():
    sheet=Sheet('Data',False,(('Name','Attendance','Quiz'),('Name','Attendance','Quiz'),('A',90,80)))
    assert header_candidates(sheet)==((1,2),'Ambiguous')
    assert len(make_raw_table(sheet,2).frame)==1

@pytest.mark.parametrize('alias',['Roll No','Roll Number','Enrollment No','Student ID','roll_no','Enrollment No.'])
def test_id_aliases(alias):
    table=raw(headers=(alias,'Name','Attendance','Quiz'))
    assert suggest_columns(table)['student_id'].selected==alias

@pytest.mark.parametrize('alias',['Name','Student Name'])
def test_name_aliases(alias):
    assert suggest_columns(raw(headers=('Roll No',alias,'Attendance','Quiz')))['student_name'].selected==alias

@pytest.mark.parametrize('alias',['Attendance','Attendance %','Attendance Percentage','attendance_percent'])
def test_attendance_aliases(alias):
    assert suggest_columns(raw(headers=('Roll No','Name',alias,'Quiz')))['attendance'].selected==alias

@pytest.mark.parametrize('alias',['Mid Sem','Quiz','End Sem','Assignment','Total','marks','final'])
def test_mark_suggestions(alias):
    assert suggest_columns(raw(headers=('Roll No','Name','Attendance',alias)))['assessments'].candidates==(alias,)


def test_mapping_ambiguity_requires_choice():
    table=raw([('A','A',90,90,80)],('Name','Student Name','Attendance','Attendance %','Quiz'))
    s=suggest_columns(table)
    assert s['student_name'].confidence=='Ambiguous' and s['student_name'].selected is None
    assert s['attendance'].confidence=='Ambiguous' and s['attendance'].selected is None
    with pytest.raises(StudentDataError):normalize_dataset(table,Mapping(None,None,None,(Assessment('Quiz','Quiz'),)))
    assert len(normalize_dataset(table,Mapping(None,'Name','Attendance',(Assessment('Quiz','Quiz'),))).frame)==1

@pytest.mark.parametrize('label,expected',[('Quiz (20)',20),('Mid Sem out of 40',40),('Quiz',None),('Quiz 1',None),('Quiz (2.5)',2.5)])
def test_maximum_suggestion(label,expected):assert suggested_maximum(label)==expected

@pytest.mark.parametrize('value,scale,expected',[(75,'auto',75),('75%','auto',75),(.75,'auto',75),(0,'auto',0),(100,'auto',100),(1.,'fraction',100),(1.,'percentage',1),(.75,'percentage',.75),(' 73% ','percentage',73)])
def test_attendance_values(value,scale,expected):assert normalize_attendance(value,scale)==expected

@pytest.mark.parametrize('value,scale',[(-1,'auto'),(101,'auto'),(1.,'auto'),(1.5,'auto'),('75%%','auto'),('abc','auto'),(1.5,'fraction'),('75%','fraction'),(float('inf'),'auto'),(True,'auto')])
def test_invalid_attendance_values(value,scale):
    with pytest.raises(StudentDataError):normalize_attendance(value,scale)

@pytest.mark.parametrize('value',['',None,'NA','N/A','null','-','NaN'])
def test_missing_never_zero(value):
    assert math.isnan(normalize_attendance(value))
    assert math.isnan(normalize_mark(value,Assessment('Quiz','Quiz')))

@pytest.mark.parametrize('value,assessment,expected',[(16,Assessment('x','x','raw',20),80),('16/20',Assessment('x','x','raw',20),80),('80%',Assessment('x','x'),80),('0',Assessment('x','x'),0),('20',Assessment('x','x','raw',20),100),(' 80 ',Assessment('x','x'),80)])
def test_mark_values(value,assessment,expected):assert normalize_mark(value,assessment)==expected

@pytest.mark.parametrize('value,assessment',[('23/20',Assessment('x','x','raw',20)),(-1,Assessment('x','x')),('eighty',Assessment('x','x')),('18/abc',Assessment('x','x','raw',20)),(101,Assessment('x','x')),('80%',Assessment('x','x','raw',20)),('16/40',Assessment('x','x','raw',20)),('16/20',Assessment('x','x')),('75%%',Assessment('x','x'))])
def test_invalid_marks(value,assessment):
    with pytest.raises(StudentDataError):normalize_mark(value,assessment)


def test_incomplete_and_invalid_rows_reported_independently():
    table=raw([('001','A',90,80),('002','B',None,80),('003','C',80,None),('004','D',145,'eighty')])
    dataset=normalize_dataset(table,mapping());assert len(dataset.frame)==3
    assert dataset.summary['valid_rows']==1 and dataset.summary['incomplete_rows']==2 and dataset.summary['invalid_rows']==1
    assert dataset.summary['invalid_attendance']==1 and dataset.summary['invalid_marks']==1
    assert dataset.summary['missing_values']==2
    result,summary=analyze_dataset(dataset)
    assert result.status.tolist()==['Normal','Incomplete','Incomplete'] and summary['incomplete_data']==2
    assert result.student_id.tolist()==['001','002','003']
    assert dataset.validation_frame().severity.tolist()==['WARNING','WARNING','ERROR','ERROR']
    assert dataset.validation_frame().row.tolist()==[3,4,5,5]


def test_partial_average_preserved():
    table=raw([('001','A',90,80,None)],('Roll No','Name','Attendance','Quiz','Mid Sem'))
    dataset=normalize_dataset(table,mapping(assessments=(Assessment('Quiz','Quiz'),Assessment('Mid Sem','Mid Sem'))))
    result,_=analyze_dataset(dataset)
    assert result.average_marks.iloc[0]==80 and result.incomplete_data.iloc[0]


def test_duplicate_ids_all_quarantined_no_merging():
    data=normalize_dataset(raw([('001','A',90,80),('001','A',90,80),('002','B',90,80)]),mapping())
    assert data.summary['duplicate_ids']==2 and data.frame.student_id.tolist()==['002']
    assert data.row_states.state.tolist()==['invalid','invalid','valid']


def test_duplicate_names_distinct_ids_retained():
    data=normalize_dataset(raw([('001','Same',90,80),('002','Same',90,30)]),mapping())
    result,_=analyze_dataset(data)
    assert result.student_id.tolist()==['001','002'] and result.status.tolist()==['Normal','Concern']


def test_name_only_duplicate_identity_reviewed():
    data=normalize_dataset(raw([('001','Same',90,80),('002','Same',90,30),('003','Unique',90,70)]),mapping(student_id=None))
    assert data.frame.student_name.tolist()==['Unique']
    assert sum(i.code=='duplicate_identity' for i in data.issues)==2


def test_missing_and_mixed_id_rejected():
    data=normalize_dataset(raw([('001','A',90,80),('', 'B',90,80)]),mapping())
    assert data.frame.student_id.tolist()==['001']
    assert any(i.code=='missing_id' and i.severity=='ERROR' for i in data.issues)


def test_id_only_display_is_id():
    data=normalize_dataset(raw(),mapping(student_name=None))
    assert data.frame.student_name.tolist()==['001']


def test_summary_rows_excluded_conservatively():
    data=normalize_dataset(raw([('', 'Class Average',90,80),('001','Total',90,80),('002','A',90,80)]),mapping())
    assert data.summary['excluded_rows']==1 and data.frame.student_name.tolist()==['Total','A']

@pytest.mark.parametrize('changes',[dict(student_id=None,student_name=None),dict(attendance=None),dict(assessments=()),dict(assessments=(Assessment('Quiz','attendance'),)),dict(assessments=(Assessment('Quiz','Quiz','raw',None),)),dict(attendance='Name'),dict(student_id='Missing'),dict(assessments=(Assessment('Quiz','Quiz'),Assessment('Quiz','Quiz2')))])
def test_invalid_mapping_controlled(changes):
    with pytest.raises(StudentDataError):normalize_dataset(raw(),mapping(**changes))


def test_equivalent_source_schemas():
    first=normalize_dataset(raw([('001','Synthetic',80,80)]),mapping()).frame
    table=raw([('001','Synthetic','.8','16/20')],('Enrollment No','Student Name','Attendance %','Mid Sem'))
    second=normalize_dataset(table,Mapping('Enrollment No','Student Name','Attendance %',(Assessment('Mid Sem','Quiz','raw',20),))).frame
    pd.testing.assert_frame_equal(first,second)


def test_trend_adapter_and_thresholds():
    table=raw([('001','Same',90,90,1),('001','Same',60,60,2),('002','Same',90,80,1),('002','Same',90,80,2)],('Roll No','Name','Attendance','Quiz','assessment_number'))
    dataset=normalize_dataset(table,mapping(assessment_number='assessment_number'))
    result,summary=analyze_dataset(dataset,Thresholds(marks_decline=20))
    assert summary['mode']=='trend' and result.student_id.tolist()==['001','002']
    assert result.marks_trend.tolist()==['Declining','Stable']
    assert result.attendance_concern.tolist()==[True,False]


def test_no_trend_claim_from_single_point():
    table=raw([('001','A',90,80,1)],('Roll No','Name','Attendance','Quiz','assessment_number'))
    result,_=analyze_dataset(normalize_dataset(table,mapping(assessment_number='assessment_number')))
    assert result.marks_trend.tolist()==['Not enough data'] and result.incomplete_data.iloc[0]


def test_export_filters_search_unicode_and_internal_exclusion():
    result,_=analyze_dataset(normalize_dataset(raw([('001','学生',90,80),('002','学生',60,30),('003','Other',None,80)]),mapping()))
    assert len(filter_students(result,'Concern'))==1
    assert len(filter_students(result,'Attendance concern'))==1
    assert len(filter_students(result,'Academic concern'))==1
    assert len(filter_students(result,'Incomplete'))==1
    shown=filter_students(result,'All','002');assert shown.student_id.tolist()==['002']
    exported=pd.read_csv(io.BytesIO(result_csv_bytes(shown)),dtype={'student_id':str})
    assert exported.student_name.tolist()==['学生'] and exported.student_id.tolist()==['002']
    assert not {'analysis_id','concern_reasons','missing_fields'}.intersection(exported.columns)
    assert 'Incomplete' in result_csv_bytes(filter_students(result,'Incomplete')).decode('utf-8-sig')


def test_csv_injection_escape():
    data=pd.DataFrame({'student_name':['=HYPERLINK("bad")','@danger','+danger','-danger'], 'marks':[0,10,20,30]})
    text=safe_csv_bytes(data).decode('utf-8-sig')
    assert "'@danger" in text and "'+danger" in text and "'-danger" in text
    assert data.student_name.iloc[0].startswith('=') # does not mutate session data


def test_validation_export_reason_codes():
    dataset=normalize_dataset(raw([('001','A',None,'bad')]),mapping())
    exported=pd.read_csv(io.BytesIO(validation_csv_bytes(dataset)))
    assert exported.code.tolist()==['missing_attendance','invalid_mark'] and exported.row.tolist()==[2,2]


def test_clear_state_only_student_workflow():
    state={'student_raw':object(),'student_sheet':'A','student_map_id':'x','student_validation':object(),'student_results':object(),'student_upload_0':b'private','studio_versions':'unrelated','navigation':'Student Data Hub','batch_letters':['old private']}
    clear_student_data(state)
    assert state=={'student_epoch':1,'studio_versions':'unrelated','navigation':'Student Data Hub'}


def test_no_ai_no_storage(monkeypatch,tmp_path):
    import ai_provider,db,vector_store
    monkeypatch.chdir(tmp_path)
    forbidden=Mock(side_effect=AssertionError('No student storage or AI'))
    monkeypatch.setattr(ai_provider,'generate_chat',forbidden)
    monkeypatch.setattr(db,'material_connection',forbidden)
    monkeypatch.setattr(db,'log_activity',forbidden)
    monkeypatch.setattr(vector_store,'get_collection',forbidden)
    book=parse_student_file(workbook_bytes(),'x.xlsx')
    dataset=normalize_dataset(make_raw_table(book.sheets[0],1),mapping(assessments=(Assessment('Quiz (20)','Quiz','raw',20),)))
    result,_=analyze_dataset(dataset);result_csv_bytes(result);validation_csv_bytes(dataset)
    forbidden.assert_not_called();assert list(tmp_path.iterdir())==[]


def test_csv_title_rows_and_blank_leading_rows():
    sheet=parse_student_file(b'\nSynthetic class results\n\nRoll No,Name,Attendance,Quiz\n001,A,90,80\n','x.csv').sheets[0]
    assert header_candidates(sheet)[0]==(4,)
    assert make_raw_table(sheet,4).source_rows==(5,)


def test_invalid_latest_history_blocks_identity():
    table=raw([('001','A',90,80,1),('001','A',90,'bad',2),('002','B',90,70,1),('002','B',90,80,2)],('Roll No','Name','Attendance','Quiz','assessment_number'))
    data=normalize_dataset(table,mapping(assessment_number='assessment_number'))
    assert data.frame.student_id.tolist()==['002','002']
    assert any(i.code=='invalid_history' and i.row==2 for i in data.issues)
    result,_=analyze_dataset(data)
    assert result.student_id.tolist()==['002']


def test_history_duplicate_assessments_and_inconsistent_names():
    table=raw([('001','A',90,80,1),('001','A',90,80,1),('002','B',90,80,1),('002','Changed',90,80,2)],('Roll No','Name','Attendance','Quiz','assessment_number'))
    data=normalize_dataset(table,mapping(assessment_number='assessment_number'))
    assert data.frame.empty
    assert {'duplicate_id','inconsistent_name'}.issubset({i.code for i in data.issues})
    with pytest.raises(StudentDataError,match='eligible'):analyze_dataset(data)


def test_missing_history_points_remain_incomplete():
    table=raw([('001','A',90,80,1),('001','A',None,None,2)],('Roll No','Name','Attendance','Quiz','assessment_number'))
    result,summary=analyze_dataset(normalize_dataset(table,mapping(assessment_number='assessment_number')))
    assert result.incomplete_data.iloc[0] and result.marks_trend.iloc[0]=='Insufficient data'
    assert pd.isna(result.latest_marks.iloc[0]) and summary['marks_available_count']==0


def test_explicit_excel_safety_options(monkeypatch):
    original=openpyxl.load_workbook
    spy=Mock(wraps=original);monkeypatch.setattr(openpyxl,'load_workbook',spy)
    parse_student_file(workbook_bytes(),'x.xlsx')
    assert spy.call_args.kwargs==dict(read_only=True,data_only=False,keep_links=False)


def test_hardened_xml_required(monkeypatch):
    import openpyxl.xml
    monkeypatch.setattr(openpyxl.xml,'DEFUSEDXML',False)
    with pytest.raises(StudentDataError,match='Safe Excel'):parse_student_file(workbook_bytes(),'x.xlsx')


def test_bounded_csv_rows_columns_cells_text(monkeypatch):
    monkeypatch.setattr(ingestion,'MAX_COLUMNS',2)
    with pytest.raises(StudentDataError,match='columns'):parse_student_file(b'A,B,C\n1,2,3\n','x.csv')
    monkeypatch.setattr(ingestion,'MAX_COLUMNS',100);monkeypatch.setattr(ingestion,'MAX_CELL_CHARS',2)
    with pytest.raises(StudentDataError,match='length'):parse_student_file(b'A,B\nlong,2\n','x.csv')


def test_identity_formula_rejected_and_trim_info():
    data=normalize_dataset(raw([('=1+1','A',90,80),(' 002 ',' B ',90,80)]),mapping())
    assert data.frame.student_id.tolist()==['002']
    assert any(i.code=='formula_identity' for i in data.issues)
    assert sum(i.code=='identity_trimmed' for i in data.issues)==2


def test_genuine_zero_mark_is_not_missing():
    data=normalize_dataset(raw([('001','A',90,0)]),mapping())
    result,_=analyze_dataset(data)
    assert result.average_marks.iloc[0]==0 and result.needs_support.iloc[0] and not result.incomplete_data.iloc[0]


def test_no_available_marks_average_is_unavailable():
    result,summary=analyze_dataset(normalize_dataset(raw([('001','A',90,None)]),mapping()))
    assert pd.isna(summary['class_average_marks']) and summary['marks_available_count']==0
    assert result.status.iloc[0]=='Incomplete' and not result.needs_support.iloc[0]


def test_public_concern_labels_distinguish_attendance_and_marks():
    result,_=analyze_dataset(normalize_dataset(raw([('001','A',60,80)]),mapping()))
    shown=public_results(result)
    assert shown.support_concern.iloc[0] and shown.attendance_concern.iloc[0]
    assert not shown.academic_concern.iloc[0]
