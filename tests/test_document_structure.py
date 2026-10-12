"""Generic template structure and AI prose boundary; synthetic provider mocks."""
import io,json
from unittest.mock import Mock
import pytest
import ai_provider
from document_models import CATALOG,DocumentRequest,DocumentError
from document_studio import generate_draft,refine_draft
from document_templates import structure_profile
from document_export import visible_blocks,document_docx_bytes,document_pdf_bytes
from tests.test_document_safety import INSTRUCTION,BODY

@pytest.mark.parametrize('kind',list(CATALOG))
def test_all_types_prose_only(monkeypatch,kind):
 profile=structure_profile(kind)
 data={'body':list(BODY),'subject':'Request for consideration'}
 if profile.family in ('letter','email'):data.update(salutation='Dear Dean,',closing='Yours sincerely,')
 provider=Mock(return_value=json.dumps(data));monkeypatch.setattr(ai_provider,'generate_chat',provider)
 result=generate_draft(DocumentRequest(kind,INSTRUCTION),approved_preferences=())
 assert result.body==BODY
 assert all(getattr(result,k)=='' for k in profile.optional_metadata)
 assert not any(k in profile.optional_metadata for k,t in visible_blocks(result,'standard_academic'))
 schema=provider.call_args.kwargs.get('response_format')
 if schema:assert schema['required']==['body']
 assert 'Omit these metadata keys' in provider.call_args.kwargs['messages'][0]['content']


def test_supplied_facts_refinement_and_exports(monkeypatch):
 provider=Mock(side_effect=[json.dumps({'body':['The proposed budget is INR 2500.'],'salutation':'Dear Dean,','closing':'Yours sincerely,'}),json.dumps({'body':['I respectfully request consideration of the proposed budget of INR 2500.']})]);monkeypatch.setattr(ai_provider,'generate_chat',provider)
 original=generate_draft(DocumentRequest('official_letter',INSTRUCTION,date='10 October 2026',additional_context='Budget INR 2500'),approved_preferences=())
 refined=refine_draft(original,'Make this more formal.')
 assert refined.date==original.date=='10 October 2026'
 assert refined.recipient==refined.sender==refined.reference_number==refined.signature==''
 from docx import Document
 text='\n'.join(p.text for p in Document(io.BytesIO(document_docx_bytes(refined))).paragraphs)
 assert '10 October 2026' in text and 'Ref:' not in text and '[date]' not in text
 import pdfplumber
 with pdfplumber.open(io.BytesIO(document_pdf_bytes(refined))) as pdf:
  text='\n'.join(p.extract_text() or '' for p in pdf.pages)
 assert '10 October 2026' in text and 'Ref:' not in text

@pytest.mark.parametrize('body',['Scheduled on [specific date].','All expenses will be covered by the university annual budget.'])
def test_prose_safety_and_two_attempt_bound(monkeypatch,body):
 provider=Mock(return_value=json.dumps({'body':[body]}));monkeypatch.setattr(ai_provider,'generate_chat',provider)
 with pytest.raises(DocumentError):generate_draft(DocumentRequest('official_letter',INSTRUCTION),approved_preferences=())
 assert provider.call_count==2


def test_blank_metadata_exports(monkeypatch):
 provider=Mock(return_value=json.dumps({'body':list(BODY),'salutation':'Dear Dean,','subject':'Request for approval','closing':'Yours sincerely,'}));monkeypatch.setattr(ai_provider,'generate_chat',provider)
 result=generate_draft(DocumentRequest('official_letter',INSTRUCTION),approved_preferences=())
 from docx import Document
 docx_text='\n'.join(p.text for p in Document(io.BytesIO(document_docx_bytes(result))).paragraphs)
 import pdfplumber
 with pdfplumber.open(io.BytesIO(document_pdf_bytes(result))) as pdf:
  pdf_text='\n'.join(p.extract_text() or '' for p in pdf.pages)
 for text in (docx_text,pdf_text):
  assert 'Dear Dean,' in text
  assert all(label not in text for label in ('Date:','Ref:','Recipient:','Sender:','[specific date]'))
 assert result.recipient==result.sender==result.signature==''
